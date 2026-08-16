"""中国监考模块 Excel 模板与批量导入。"""
from datetime import date, datetime
from hashlib import sha1
from io import BytesIO
from typing import Optional

from fastapi import APIRouter, Depends, File, HTTPException, Query, UploadFile
from fastapi.responses import StreamingResponse
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from sqlalchemy.orm import Session

from auth_utils import require_admin
from dependencies import get_db
from invigilation_defaults import seed_defaults
from invigilation_models import (
    InvigilationDutyRequirement,
    InvigilationDutyRole,
    InvigilationExamBatch,
    InvigilationExamRoom,
    InvigilationExamSession,
    InvigilationLocation,
    InvigilationTeacherProfile,
    InvigilationTeacherRoleQualification,
    InvigilationTeacherUnavailability,
)

router = APIRouter(prefix="/api/invigilation/import", tags=["invigilation-import"])

TRUE_VALUES = {"1", "true", "yes", "y", "是", "√", "允许", "可"}
FALSE_VALUES = {"0", "false", "no", "n", "否", "×", "不允许", "不可"}
ROLE_HEADER_MAP = {
    "甲监资格": "PRIMARY",
    "乙监资格": "SECONDARY",
    "流动监考资格": "ROVING",
    "巡考资格": "INSPECTOR",
}


def _text(value) -> Optional[str]:
    if value is None:
        return None
    value = str(value).strip()
    return value or None


def _int(value, default=None):
    if value is None or str(value).strip() == "":
        return default
    try:
        return int(float(value))
    except (TypeError, ValueError):
        raise ValueError(f"无法转换为整数: {value}")


def _bool(value, default=None):
    if value is None or str(value).strip() == "":
        return default
    token = str(value).strip().lower()
    if token in TRUE_VALUES:
        return True
    if token in FALSE_VALUES:
        return False
    raise ValueError(f"无法识别的是/否值: {value}")


def _date(value) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = _text(value)
    if not text:
        raise ValueError("日期不能为空")
    return date.fromisoformat(text.replace("/", "-").replace(".", "-"))


def _rows(ws):
    headers = [_text(c.value) for c in ws[1]]
    for idx, cells in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if not any(v is not None and str(v).strip() != "" for v in cells):
            continue
        yield idx, {headers[i]: cells[i] for i in range(min(len(headers), len(cells))) if headers[i]}


def _require_headers(ws, required):
    headers = {_text(c.value) for c in ws[1]}
    missing = [h for h in required if h not in headers]
    if missing:
        raise HTTPException(status_code=400, detail=f"工作表 {ws.title} 缺少列: {', '.join(missing)}")


def _style_header(ws):
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.freeze_panes = "A2"
    for column in ws.columns:
        letter = column[0].column_letter
        width = min(28, max(12, max(len(str(c.value or "")) for c in column) + 2))
        ws.column_dimensions[letter].width = width


def _xlsx_response(wb: Workbook, filename: str):
    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.get("/teacher-template")
def teacher_template(_=Depends(require_admin)):
    wb = Workbook()
    ws = wb.active
    ws.title = "人员"
    ws.append([
        "工号", "姓名", "部门", "甲监资格", "乙监资格", "流动监考资格", "巡考资格",
        "最大总场次", "每日最大场次", "启用", "备注",
    ])
    ws.append(["10001", "张三", "计算机学院", "是", "是", "否", "否", 6, 2, "是", "示例数据，导入前可删除"])
    ws.append(["10002", "李四", "教务处", "否", "否", "否", "是", 4, 1, "是", None])
    _style_header(ws)

    unavailable = wb.create_sheet("不可用时间")
    unavailable.append(["工号", "姓名", "日期", "时段代码", "原因"])
    unavailable.append(["10001", "张三", "2026-12-25", "AM", "已有教学任务"])
    unavailable.append(["10002", "李四", "2026-12-26", None, "整天不可用"])
    _style_header(unavailable)
    return _xlsx_response(wb, "invigilation_teachers_template.xlsx")


@router.get("/exam-template")
def exam_template(_=Depends(require_admin)):
    wb = Workbook()
    ws = wb.active
    ws.title = "考试安排"
    ws.append([
        "日期", "时段代码", "时段名称", "开始时间", "结束时间", "课程代码", "课程名称", "班级",
        "校区", "考区", "楼栋", "楼层", "考场", "考场容量", "考生人数", "任课教师", "部门",
        "甲监人数", "乙监人数", "流动监考人数", "巡考人数", "备注",
    ])
    ws.append([
        "2026-12-25", "AM", "上午", "09:00", "11:00", "MATH001", "高等数学", "2025级1班",
        "东校区", "第一考区", "第一教学楼", "1层", "A101", 40, 36, "王老师", "数学学院",
        1, 1, 2, 2, "流动人数按楼层、巡考人数按考区；同范围重复行会覆盖为相同值",
    ])
    ws.append([
        "2026-12-25", "AM", "上午", "09:00", "11:00", "MATH001", "高等数学", "2025级2班",
        "东校区", "第一考区", "第一教学楼", "1层", "A102", 40, 38, "王老师", "数学学院",
        1, 1, 2, 2, None,
    ])
    _style_header(ws)
    return _xlsx_response(wb, "invigilation_exams_template.xlsx")


def _find_teacher(db: Session, employee_no: Optional[str], name: Optional[str]):
    if employee_no:
        row = db.query(InvigilationTeacherProfile).filter(
            InvigilationTeacherProfile.employee_no == employee_no
        ).first()
        if row:
            return row
    if name:
        return db.query(InvigilationTeacherProfile).filter(
            InvigilationTeacherProfile.professor_name == name
        ).first()
    return None


@router.post("/teachers")
def import_teachers(file: UploadFile = File(...), db: Session = Depends(get_db), _=Depends(require_admin)):
    seed_defaults(db)
    if not file.filename or not file.filename.lower().endswith(".xlsx"):
        raise HTTPException(status_code=400, detail="仅支持 .xlsx 文件")
    try:
        wb = load_workbook(file.file, data_only=True)
    except Exception as exc:
        raise HTTPException(status_code=400, detail=f"无法读取 Excel: {exc}")
    if "人员" not in wb.sheetnames:
        raise HTTPException(status_code=400, detail="缺少“人员”工作表")

    ws = wb["人员"]
    _require_headers(ws, ["姓名"])
    roles = {r.code: r for r in db.query(InvigilationDutyRole).all()}
    created = updated = qualification_changes = 0
    errors = []

    for row_no, values in _rows(ws):
        row_created = row_updated = row_q = 0
        try:
            with db.begin_nested():
                name = _text(values.get("姓名"))
                if not name:
                    raise ValueError("姓名不能为空")
                employee_no = _text(values.get("工号"))
                profile = _find_teacher(db, employee_no, name)
                is_new = profile is None
                if is_new:
                    profile = InvigilationTeacherProfile(professor_name=name)
                    db.add(profile)
                    db.flush()
                    row_created = 1
                else:
                    row_updated = 1

                profile.professor_name = name
                profile.employee_no = employee_no
                profile.department = _text(values.get("部门"))
                profile.max_total_duties = _int(values.get("最大总场次"), None)
                profile.max_daily_duties = _int(values.get("每日最大场次"), 2)
                profile.active = _bool(values.get("启用"), True)
                profile.notes = _text(values.get("备注"))

                for header, role_code in ROLE_HEADER_MAP.items():
                    role = roles.get(role_code)
                    if not role:
                        continue
                    default = role_code in {"PRIMARY", "SECONDARY"} if is_new else None
                    eligible = _bool(values.get(header), default)
                    if eligible is None:
                        continue
                    q = db.query(InvigilationTeacherRoleQualification).filter(
                        InvigilationTeacherRoleQualification.teacher_profile_id == profile.id,
                        InvigilationTeacherRoleQualification.duty_role_id == role.id,
                    ).first()
                    if not q:
                        db.add(InvigilationTeacherRoleQualification(
                            teacher_profile_id=profile.id,
                            duty_role_id=role.id,
                            eligible=eligible,
                        ))
                    else:
                        q.eligible = eligible
                    row_q += 1
                db.flush()
            created += row_created
            updated += row_updated
            qualification_changes += row_q
        except Exception as exc:
            errors.append({"sheet": "人员", "row": row_no, "message": str(exc)})

    unavailability_count = 0
    if "不可用时间" in wb.sheetnames:
        uw = wb["不可用时间"]
        _require_headers(uw, ["日期"])
        for row_no, values in _rows(uw):
            try:
                with db.begin_nested():
                    profile = _find_teacher(
                        db, _text(values.get("工号")), _text(values.get("姓名"))
                    )
                    if not profile:
                        raise ValueError("找不到对应人员，请填写正确工号或姓名")
                    day = _date(values.get("日期"))
                    slot_code = _text(values.get("时段代码"))
                    query = db.query(InvigilationTeacherUnavailability).filter(
                        InvigilationTeacherUnavailability.teacher_profile_id == profile.id,
                        InvigilationTeacherUnavailability.date == day,
                    )
                    existing = (
                        query.filter(InvigilationTeacherUnavailability.slot_code.is_(None)).first()
                        if slot_code is None
                        else query.filter(InvigilationTeacherUnavailability.slot_code == slot_code).first()
                    )
                    if not existing:
                        existing = InvigilationTeacherUnavailability(
                            teacher_profile_id=profile.id, date=day, slot_code=slot_code
                        )
                        db.add(existing)
                    existing.reason = _text(values.get("原因"))
                    db.flush()
                unavailability_count += 1
            except Exception as exc:
                errors.append({"sheet": "不可用时间", "row": row_no, "message": str(exc)})

    db.commit()
    return {
        "created": created,
        "updated": updated,
        "qualification_changes": qualification_changes,
        "unavailability_rows": unavailability_count,
        "errors": errors,
    }


def _location_code(location_type: str, path: str) -> str:
    return f"AUTO-{location_type}-{sha1(path.encode('utf-8')).hexdigest()[:14]}"


def _get_or_create_location(db: Session, name: Optional[str], location_type: str,
                            parent: Optional[InvigilationLocation], path: str,
                            capacity: Optional[int] = None):
    if not name:
        return parent
    query = db.query(InvigilationLocation).filter(
        InvigilationLocation.name == name,
        InvigilationLocation.location_type == location_type,
    )
    query = (
        query.filter(InvigilationLocation.parent_id.is_(None))
        if parent is None
        else query.filter(InvigilationLocation.parent_id == parent.id)
    )
    row = query.first()
    if not row:
        row = InvigilationLocation(
            code=_location_code(location_type, path),
            name=name,
            location_type=location_type,
            parent_id=parent.id if parent else None,
            capacity=capacity,
        )
        db.add(row)
        db.flush()
    elif capacity is not None:
        row.capacity = capacity
    return row


def _upsert_requirement(db: Session, session_id: int, location_id: int,
                        role: InvigilationDutyRole, count: Optional[int]):
    if count is None:
        return 0
    row = db.query(InvigilationDutyRequirement).filter(
        InvigilationDutyRequirement.session_id == session_id,
        InvigilationDutyRequirement.location_id == location_id,
        InvigilationDutyRequirement.duty_role_id == role.id,
    ).first()
    if not row:
        row = InvigilationDutyRequirement(
            session_id=session_id,
            location_id=location_id,
            duty_role_id=role.id,
        )
        db.add(row)
    row.required_count = count
    return 1


@router.post("/exams")
def import_exams(
    file: UploadFile = File(...),
    batch_id: int = Query(..., ge=1),
    db: Session = Depends(get_db),
    _=Depends(require_admin),
):
    seed_defaults(db)
    if not db.get(InvigilationExamBatch, batch_id):
        raise HTTPException(status_code=404, detail="考试批次不存在，请先创建考试批次")
    if not file.filename or not file.filename.lower().endswith(".xlsx"):
        raise HTTPException(status_code=400, detail="仅支持 .xlsx 文件")
    try:
        wb = load_workbook(file.file, data_only=True)
    except Exception as exc:
        raise HTTPException(status_code=400, detail=f"无法读取 Excel: {exc}")
    if "考试安排" not in wb.sheetnames:
        raise HTTPException(status_code=400, detail="缺少“考试安排”工作表")
    ws = wb["考试安排"]
    _require_headers(ws, ["日期", "时段代码", "课程名称", "考场"])

    roles = {r.code: r for r in db.query(InvigilationDutyRole).all()}
    created_sessions = created_exam_rooms = updated_exam_rooms = requirement_changes = 0
    errors = []

    for row_no, values in _rows(ws):
        row_session = row_created_room = row_updated_room = row_req = 0
        try:
            with db.begin_nested():
                exam_date = _date(values.get("日期"))
                slot_code = _text(values.get("时段代码"))
                course_name = _text(values.get("课程名称"))
                room_name = _text(values.get("考场"))
                if not slot_code or not course_name or not room_name:
                    raise ValueError("时段代码、课程名称、考场不能为空")

                session = db.query(InvigilationExamSession).filter(
                    InvigilationExamSession.batch_id == batch_id,
                    InvigilationExamSession.date == exam_date,
                    InvigilationExamSession.slot_code == slot_code,
                ).first()
                if not session:
                    session = InvigilationExamSession(
                        batch_id=batch_id,
                        date=exam_date,
                        slot_code=slot_code,
                        slot_name=_text(values.get("时段名称")),
                        start_time=_text(values.get("开始时间")),
                        end_time=_text(values.get("结束时间")),
                    )
                    db.add(session)
                    db.flush()
                    row_session = 1

                names = [
                    (_text(values.get("校区")), "CAMPUS"),
                    (_text(values.get("考区")), "AREA"),
                    (_text(values.get("楼栋")), "BUILDING"),
                    (_text(values.get("楼层")), "FLOOR"),
                    (room_name, "ROOM"),
                ]
                parent = None
                path_parts = []
                locations = {}
                for loc_name, loc_type in names:
                    if not loc_name:
                        continue
                    path_parts.append(loc_name)
                    capacity = _int(values.get("考场容量"), None) if loc_type == "ROOM" else None
                    parent = _get_or_create_location(
                        db, loc_name, loc_type, parent, "/".join(path_parts), capacity
                    )
                    locations[loc_type] = parent

                room = locations.get("ROOM")
                if not room:
                    raise ValueError("无法建立考场地点")
                exam_room = db.query(InvigilationExamRoom).filter(
                    InvigilationExamRoom.session_id == session.id,
                    InvigilationExamRoom.room_location_id == room.id,
                ).first()
                if not exam_room:
                    exam_room = InvigilationExamRoom(session_id=session.id, room_location_id=room.id)
                    db.add(exam_room)
                    row_created_room = 1
                else:
                    row_updated_room = 1
                exam_room.course_code = _text(values.get("课程代码"))
                exam_room.course_name = course_name
                exam_room.group_names = _text(values.get("班级"))
                exam_room.candidate_count = _int(values.get("考生人数"), None)
                exam_room.department = _text(values.get("部门"))
                exam_room.course_teacher_names = _text(values.get("任课教师"))
                exam_room.notes = _text(values.get("备注"))
                db.flush()

                row_req += _upsert_requirement(
                    db, session.id, room.id, roles["PRIMARY"],
                    _int(values.get("甲监人数"), roles["PRIMARY"].default_required),
                )
                row_req += _upsert_requirement(
                    db, session.id, room.id, roles["SECONDARY"],
                    _int(values.get("乙监人数"), roles["SECONDARY"].default_required),
                )

                roving_count = _int(values.get("流动监考人数"), None)
                if roving_count is not None:
                    floor = locations.get("FLOOR")
                    if not floor:
                        raise ValueError("填写了流动监考人数时必须提供楼层")
                    row_req += _upsert_requirement(db, session.id, floor.id, roles["ROVING"], roving_count)

                inspector_count = _int(values.get("巡考人数"), None)
                if inspector_count is not None:
                    area = locations.get("AREA")
                    if not area:
                        raise ValueError("填写了巡考人数时必须提供考区")
                    row_req += _upsert_requirement(db, session.id, area.id, roles["INSPECTOR"], inspector_count)
                db.flush()

            created_sessions += row_session
            created_exam_rooms += row_created_room
            updated_exam_rooms += row_updated_room
            requirement_changes += row_req
        except Exception as exc:
            errors.append({"sheet": "考试安排", "row": row_no, "message": str(exc)})

    db.commit()
    return {
        "batch_id": batch_id,
        "created_sessions": created_sessions,
        "created_exam_rooms": created_exam_rooms,
        "updated_exam_rooms": updated_exam_rooms,
        "requirement_changes": requirement_changes,
        "errors": errors,
    }
