"""中国监考模块报表导出：总监考表、个人明细、工作量统计。"""
from collections import defaultdict
from io import BytesIO
from urllib.parse import quote

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from sqlalchemy.orm import Session

from auth_utils import require_admin
from dependencies import get_db
from invigilation_models import (
    InvigilationAssignment,
    InvigilationDutyRole,
    InvigilationExamBatch,
    InvigilationExamSession,
    InvigilationLocation,
    InvigilationTeacherProfile,
)

router = APIRouter(prefix="/api/invigilation/export", tags=["invigilation-export"])

ROLE_ORDER = {"PRIMARY": 10, "SECONDARY": 20, "ROVING": 30, "INSPECTOR": 40}


def _batch_or_404(db: Session, batch_id: int):
    batch = db.get(InvigilationExamBatch, batch_id)
    if not batch:
        raise HTTPException(status_code=404, detail="考试批次不存在")
    return batch


def _assignment_rows(db: Session, batch_id: int):
    rows = (
        db.query(
            InvigilationAssignment,
            InvigilationExamSession,
            InvigilationLocation,
            InvigilationDutyRole,
            InvigilationTeacherProfile,
        )
        .join(InvigilationExamSession, InvigilationAssignment.session_id == InvigilationExamSession.id)
        .join(InvigilationLocation, InvigilationAssignment.location_id == InvigilationLocation.id)
        .join(InvigilationDutyRole, InvigilationAssignment.duty_role_id == InvigilationDutyRole.id)
        .join(InvigilationTeacherProfile, InvigilationAssignment.teacher_profile_id == InvigilationTeacherProfile.id)
        .filter(InvigilationExamSession.batch_id == batch_id)
        .all()
    )
    rows.sort(key=lambda item: (
        item[1].date,
        item[1].sort_order,
        item[1].slot_code,
        item[2].name,
        ROLE_ORDER.get(item[3].code, item[3].sort_order),
        item[4].professor_name,
    ))
    return rows


def _content_disposition(ascii_name: str, display_name: str) -> str:
    """RFC 5987 filename*: 兼容 Starlette latin-1 header 与 Windows 中文文件名。"""
    encoded = quote(display_name, safe="")
    return f'attachment; filename="{ascii_name}"; filename*=UTF-8\'\'{encoded}'


def _xlsx_response(workbook: Workbook, ascii_name: str, display_name: str):
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": _content_disposition(ascii_name, display_name)},
    )


def _style_sheet(ws):
    header_fill = PatternFill("solid", fgColor="DCE6F1")
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max(max_len + 2, 10), 28)


@router.get("/{batch_id}/xlsx")
def export_xlsx(batch_id: int, db: Session = Depends(get_db), _=Depends(require_admin)):
    batch = _batch_or_404(db, batch_id)
    rows = _assignment_rows(db, batch_id)
    if not rows:
        raise HTTPException(status_code=400, detail="当前批次还没有监考排班结果")

    wb = Workbook()
    total = wb.active
    total.title = "总监考表"
    total.append(["日期", "时段", "开始", "结束", "地点/考场", "岗位", "工号", "监考人员", "部门", "锁定", "来源"])

    teacher_rows = []
    role_counts = defaultdict(lambda: defaultdict(int))
    weighted = defaultdict(float)
    teacher_info = {}

    for assignment, session, location, role, teacher in rows:
        slot = session.slot_name or session.slot_code
        line = [
            session.date.isoformat(), slot, session.start_time, session.end_time,
            location.name, role.name, teacher.employee_no, teacher.professor_name,
            teacher.department, "是" if assignment.locked else "否", assignment.source,
        ]
        total.append(line)
        teacher_rows.append((teacher.professor_name, *line))
        role_counts[teacher.id][role.code] += 1
        weighted[teacher.id] += role.workload_weight
        teacher_info[teacher.id] = teacher
    _style_sheet(total)

    personal = wb.create_sheet("按教师明细")
    personal.append(["监考人员", "日期", "时段", "开始", "结束", "地点/考场", "岗位", "工号", "姓名", "部门", "锁定", "来源"])
    for row in sorted(teacher_rows, key=lambda x: (x[0], x[1], x[2], x[5], x[6])):
        personal.append(row)
    _style_sheet(personal)

    workload = wb.create_sheet("工作量统计")
    workload.append(["工号", "姓名", "部门", "甲监", "乙监", "流动监考", "巡考", "总次数", "加权工作量"])
    for teacher_id, teacher in sorted(teacher_info.items(), key=lambda item: item[1].professor_name):
        counts = role_counts[teacher_id]
        workload.append([
            teacher.employee_no,
            teacher.professor_name,
            teacher.department,
            counts["PRIMARY"],
            counts["SECONDARY"],
            counts["ROVING"],
            counts["INSPECTOR"],
            sum(counts.values()),
            round(weighted[teacher_id], 3),
        ])
    _style_sheet(workload)

    meta = wb.create_sheet("批次信息")
    meta.append(["项目", "内容"])
    meta.append(["批次名称", batch.name])
    meta.append(["学年", batch.academic_year])
    meta.append(["学期", batch.term])
    meta.append(["状态", batch.status])
    meta.append(["监考岗位总数", len(rows)])
    _style_sheet(meta)

    return _xlsx_response(
        wb,
        f"invigilation_{batch_id}.xlsx",
        f"{batch.name}_监考安排.xlsx",
    )


def _register_chinese_font():
    font_name = "STSong-Light"
    try:
        pdfmetrics.getFont(font_name)
    except KeyError:
        pdfmetrics.registerFont(UnicodeCIDFont(font_name))
    return font_name


@router.get("/{batch_id}/pdf")
def export_pdf(batch_id: int, db: Session = Depends(get_db), _=Depends(require_admin)):
    batch = _batch_or_404(db, batch_id)
    rows = _assignment_rows(db, batch_id)
    if not rows:
        raise HTTPException(status_code=400, detail="当前批次还没有监考排班结果")

    font_name = _register_chinese_font()
    output = BytesIO()
    doc = SimpleDocTemplate(
        output,
        pagesize=landscape(A4),
        rightMargin=10 * mm,
        leftMargin=10 * mm,
        topMargin=10 * mm,
        bottomMargin=10 * mm,
        title=f"{batch.name} 监考安排",
    )
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "CnTitle", parent=styles["Title"], fontName=font_name,
        fontSize=16, leading=22, alignment=TA_CENTER, spaceAfter=8,
    )
    note_style = ParagraphStyle(
        "CnNote", parent=styles["Normal"], fontName=font_name,
        fontSize=8, leading=11, alignment=TA_CENTER,
        textColor=colors.HexColor("#555555"),
    )

    story = [
        Paragraph(f"{batch.name} — 监考安排表", title_style),
        Paragraph(
            "岗位：甲监 / 乙监 / 流动监考 / 巡考。带“锁定”的安排在局部重排时保持不变。",
            note_style,
        ),
        Spacer(1, 5 * mm),
    ]
    data = [["日期", "时段", "地点/考场", "岗位", "监考人员", "工号", "部门", "锁定"]]
    for assignment, session, location, role, teacher in rows:
        data.append([
            session.date.isoformat(),
            session.slot_name or session.slot_code,
            location.name,
            role.name,
            teacher.professor_name,
            teacher.employee_no or "",
            teacher.department or "",
            "是" if assignment.locked else "",
        ])

    table = Table(data, repeatRows=1, colWidths=[27*mm, 24*mm, 38*mm, 26*mm, 31*mm, 27*mm, 44*mm, 18*mm])
    table.setStyle(TableStyle([
        ("FONTNAME", (0, 0), (-1, -1), font_name),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#DCE6F1")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#1F2937")),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#B8C2CC")),
        ("BOTTOMPADDING", (0, 0), (-1, 0), 6),
        ("TOPPADDING", (0, 0), (-1, 0), 6),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F7F9FC")]),
    ]))
    story.append(table)
    doc.build(story)
    output.seek(0)

    return StreamingResponse(
        output,
        media_type="application/pdf",
        headers={
            "Content-Disposition": _content_disposition(
                f"invigilation_{batch_id}.pdf",
                f"{batch.name}_监考安排.pdf",
            )
        },
    )
